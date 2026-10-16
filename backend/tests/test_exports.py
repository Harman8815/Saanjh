"""
Tests for utils.exports CSV/XLSX download helpers.
"""

import csv
import io

from django.test import TestCase

from utils.exports import write_csv, write_xlsx


class WriteCsvTests(TestCase):
    def test_headers_and_rows(self):
        response = write_csv(
            [
                ['Ada Lovelace', 'ada@example.com', 2],
                ['Grace Hopper', 'grace@example.com', 1],
            ],
            ['Name', 'Email', 'Guests'],
            filename='guests.csv',
        )

        self.assertEqual(response['Content-Type'], 'text/csv')
        self.assertEqual(
            response['Content-Disposition'],
            'attachment; filename="guests.csv"',
        )

        rows = list(csv.reader(io.StringIO(response.content.decode('utf-8'))))
        self.assertEqual(rows[0], ['Name', 'Email', 'Guests'])
        self.assertEqual(rows[1], ['Ada Lovelace', 'ada@example.com', '2'])
        self.assertEqual(rows[2], ['Grace Hopper', 'grace@example.com', '1'])

    def test_dict_rows_use_header_keys(self):
        response = write_csv(
            [{'first_name': 'Ada', 'email': 'ada@example.com', 'missing': None}],
            [('first_name', 'First name'), ('email', 'Email'), ('rsvp', 'RSVP')],
        )
        rows = list(csv.reader(io.StringIO(response.content.decode('utf-8'))))
        # A missing dict key renders as an empty cell, None too.
        self.assertEqual(rows[1], ['Ada', 'ada@example.com', ''])

    def test_none_values_render_empty(self):
        response = write_csv([[None, 'value']], ['A', 'B'])
        rows = list(csv.reader(io.StringIO(response.content.decode('utf-8'))))
        self.assertEqual(rows[1], ['', 'value'])

    def test_empty_rows_still_writes_header(self):
        response = write_csv([], ['Only header'])
        rows = list(csv.reader(io.StringIO(response.content.decode('utf-8'))))
        self.assertEqual(rows, [['Only header']])


class WriteXlsxTests(TestCase):
    def test_xlsx_payload(self):
        response = write_xlsx(
            [['Ada Lovelace', 2], ['Grace Hopper', None]],
            ['Name', 'Guests'],
            filename='guests.xlsx',
            sheet_name='Guests',
        )

        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertEqual(
            response['Content-Disposition'],
            'attachment; filename="guests.xlsx"',
        )
        # An xlsx file is a zip container: its first two bytes are 'PK'.
        self.assertTrue(response.content.startswith(b'PK'))

        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Guests'])
        sheet = workbook['Guests']
        self.assertEqual(sheet['A1'].value, 'Name')
        self.assertEqual(sheet['A2'].value, 'Ada Lovelace')
        self.assertEqual(sheet['B2'].value, 2)
        # None renders as an empty cell, not the string "None".
        self.assertIsNone(sheet['B3'].value)

    def test_xlsx_accepts_dict_rows(self):
        response = write_xlsx(
            [{'name': 'Ada', 'count': 3}],
            [('name', 'Name'), ('count', 'Count')],
        )
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet['A1'].value, 'Name')
        self.assertEqual(sheet['B2'].value, 3)
