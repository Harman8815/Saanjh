"""
Guest spreadsheet import: parse, validate per row, and commit.

The flow is always two-phase:

1. ``preview_import`` parses the upload and returns per-row validation plus
   flagged duplicates. The validated rows are cached server-side under a
   short-lived token (``cache_import_token``) so phase 2 never trusts rows
   that round-tripped through the client.
2. ``commit_import`` reads the cached rows and creates guests.

Rows are capped at MAX_IMPORT_ROWS and the file itself is checked by
``utils.uploads.validate_uploaded_file`` before parsing.
"""

import csv
import io
import re
import uuid

from rest_framework import serializers

from utils.uploads import validate_uploaded_file

# Template columns, as (field key, human header).
GUEST_IMPORT_HEADERS = [
    ('first_name', 'First Name'),
    ('last_name', 'Last Name'),
    ('email', 'Email'),
    ('phone', 'Phone'),
    ('relationship', 'Relationship'),
    ('rsvp_status', 'RSVP Status'),
    ('dietary_restrictions', 'Dietary Restrictions'),
    ('address', 'Address'),
    ('notes', 'Notes'),
    ('group', 'Group'),
]

MAX_IMPORT_ROWS = 5000

# How long a validated preview stays available for confirmation.
PREVIEW_CACHE_TTL = 900  # seconds

_EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')


def normalise_header(name):
    """'RSVP Status' -> 'rsvp_status' so headers match case-insensitively."""
    return re.sub(r'[^a-z0-9]+', '_', (name or '').strip().lower()).strip('_')


def cache_key_for_token(token):
    return f'guest-import-preview:{token}'


def cache_import_token(token, rows):
    from django.core.cache import cache
    cache.set(cache_key_for_token(token), rows, PREVIEW_CACHE_TTL)


def get_cached_import(token):
    from django.core.cache import cache
    return cache.get(cache_key_for_token(token))


def clear_cached_import(token):
    from django.core.cache import cache
    cache.delete(cache_key_for_token(token))


def parse_spreadsheet(upload):
    """
    Parse a CSV or XLSX upload into a list of header-keyed dicts.

    Raises DRF ValidationError for unsupported files, missing headers, or
    files over the row cap.
    """
    validate_uploaded_file(upload, 'guestlist')

    filename = (getattr(upload, 'name', '') or '').lower()
    if filename.endswith('.csv'):
        rows = _parse_csv(upload)
    elif filename.endswith('.xlsx'):
        rows = _parse_xlsx(upload)
    else:
        raise serializers.ValidationError(
            'Unsupported file type. Upload a .csv or .xlsx file.'
        )

    if not rows:
        raise serializers.ValidationError('The uploaded file has no data rows.')

    if len(rows) > MAX_IMPORT_ROWS:
        raise serializers.ValidationError(
            f'The file has {len(rows)} rows, which exceeds the '
            f'{MAX_IMPORT_ROWS}-row import limit.'
        )
    return rows


def _parse_csv(upload):
    raw = upload.read()
    if isinstance(raw, bytes):
        # utf-8-sig strips the BOM Excel adds when saving CSV.
        raw = raw.decode('utf-8-sig', errors='replace')
    reader = csv.DictReader(io.StringIO(raw))
    if not reader.fieldnames:
        raise serializers.ValidationError('The CSV file has no header row.')

    rows = []
    for line_number, row in enumerate(reader, start=2):  # header is line 1
        rows.append(_normalise_row(row, line_number))
    return rows


def _parse_xlsx(upload):
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(upload.read()), read_only=True, data_only=True)
    sheet = workbook.active
    header_iter = sheet.iter_rows(values_only=True)

    try:
        raw_headers = [str(cell).strip() if cell is not None else ''
                       for cell in next(header_iter)]
    except StopIteration:
        raise serializers.ValidationError('The spreadsheet has no header row.')

    headers = [normalise_header(header) for header in raw_headers]
    rows = []
    for line_number, values in enumerate(header_iter, start=2):
        if all(value is None or str(value).strip() == '' for value in values):
            continue  # skip fully blank rows
        row = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            value = values[index] if index < len(values) else None
            row[header] = '' if value is None else str(value).strip()
        rows.append(_normalise_row(row, line_number))
    return rows


def _normalise_row(raw_row, line_number):
    """Map raw columns to known field keys, dropping unknown columns."""
    row = {'row': line_number}
    for key, value in raw_row.items():
        normalised = normalise_header(key) if key else ''
        if normalised in {field for field, _ in GUEST_IMPORT_HEADERS}:
            row[normalised] = ('' if value is None else str(value)).strip()
    return row


def validate_rows(rows, wedding):
    """
    Validate each parsed row against the wedding's reference data.

    Unknown rsvp_status, relationship, group, or email formats are reported
    as errors; nothing is silently coerced. Returns a list of preview dicts:

    ``{row, valid, data, errors, is_duplicate}``
    """
    from .models import Guest, GuestGroup, RsvpStatus

    rsvp_statuses = {status.name.lower(): status.id for status in RsvpStatus.objects.all()}
    groups = {group.name.lower(): group.id for group in GuestGroup.objects.filter(wedding=wedding)}
    # A guest is a duplicate when an existing row shares the email, or the
    # full name. This is the flag shown in the preview; the scoring endpoint
    # (/guests/duplicates/) does the finer-grained detection.
    existing_emails = {
        email.lower() for email in
        Guest.objects.filter(wedding=wedding).exclude(email='').values_list('email', flat=True)
    }
    existing_names = {
        f"{first} {last}".lower() for first, last in
        Guest.objects.filter(wedding=wedding).values_list('first_name', 'last_name')
    }
    relationships = dict(Guest.RELATIONSHIP_CHOICES)

    previews = []
    for row in rows:
        errors = []
        data = {
            'first_name': row.get('first_name', ''),
            'last_name': row.get('last_name', ''),
            'email': row.get('email', ''),
            'phone': row.get('phone', ''),
            'address': row.get('address', ''),
            'notes': row.get('notes', ''),
            'dietary_restrictions': row.get('dietary_restrictions', ''),
        }

        # Name: at least one of first/last is required (the model allows both
        # to be blank, but an unnamed guest is useless on a seating plan).
        if not data['first_name'] and not data['last_name']:
            errors.append('A first or last name is required.')

        if data['email'] and not _EMAIL_RE.match(data['email']):
            errors.append(f'"{data["email"]}" is not a valid email address.')

        relationship = (row.get('relationship') or '').lower()
        if relationship:
            if relationship not in relationships:
                errors.append(
                    f'Unknown relationship "{row.get("relationship")}". '
                    f'Allowed: {", ".join(sorted(relationships))}.'
                )
            else:
                data['relationship'] = relationship

        rsvp_status = (row.get('rsvp_status') or '').lower()
        if rsvp_status:
            if rsvp_status not in rsvp_statuses:
                errors.append(
                    f'Unknown RSVP status "{row.get("rsvp_status")}". '
                    f'Allowed: {", ".join(sorted(rsvp_statuses))}.'
                )
            else:
                data['rsvp_status_id'] = rsvp_statuses[rsvp_status]

        group_name = row.get('group', '')
        if group_name:
            group_id = groups.get(group_name.lower())
            if group_id is None:
                errors.append(
                    f'Unknown group "{group_name}". Create it first or remove the value.'
                )
            else:
                data['group_id'] = group_id

        is_duplicate = (
            bool(data['email']) and data['email'].lower() in existing_emails
        ) or (
            bool(data['first_name'] or data['last_name'])
            and f"{data['first_name']} {data['last_name']}".strip().lower() in existing_names
        )

        previews.append({
            'row': row['row'],
            'valid': not errors,
            'data': data if not errors else None,
            'errors': errors,
            'is_duplicate': is_duplicate,
        })

    return previews


def import_preview_summary(previews):
    return {
        'total_rows': len(previews),
        'valid_rows': sum(1 for p in previews if p['valid']),
        'invalid_rows': sum(1 for p in previews if not p['valid']),
        'duplicate_rows': sum(1 for p in previews if p['is_duplicate']),
    }


def commit_valid_rows(wedding, previews):
    """
    Create guests from already-validated preview rows.

    Only rows whose ``valid`` flag is set are imported; invalid and duplicate
    rows are skipped and reported. Guests are created one by one so
    ``Guest.save()`` and its defaults run for every row.
    """
    from .models import Guest

    created = 0
    skipped = 0
    errors = []

    for preview in previews:
        data = preview.get('data')
        if not preview.get('valid') or not data:
            skipped += 1
            continue
        if preview.get('is_duplicate'):
            # Duplicates are never auto-merged on import; they are skipped so
            # the couple can resolve them in the duplicates view.
            skipped += 1
            continue
        try:
            # rsvp_status has a model-level default, so only set it when the
            # row supplied a valid status (passing None would override the
            # default with NULL and fail the NOT NULL constraint).
            guest_kwargs = dict(
                wedding=wedding,
                first_name=data.get('first_name') or '',
                last_name=data.get('last_name') or '',
                email=data.get('email') or None,
                phone=data.get('phone') or None,
                address=data.get('address') or None,
                notes=data.get('notes') or None,
                dietary_restrictions=data.get('dietary_restrictions') or None,
                relationship=data.get('relationship', 'friend'),
                group_id=data.get('group_id'),
            )
            if data.get('rsvp_status_id'):
                guest_kwargs['rsvp_status_id'] = data['rsvp_status_id']
            Guest.objects.create(**guest_kwargs)
            created += 1
        except Exception as exc:  # one bad row must not abort the batch
            skipped += 1
            errors.append(f"Row {preview['row']}: {exc}")

    return {
        'created': created,
        'skipped': skipped,
        'errors': errors,
    }


def new_import_token():
    return uuid.uuid4().hex
